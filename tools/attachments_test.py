"""Builds sample attachments (PDF with a text page and a drawing page, docx, xlsx, pptx,
csv, step, png) and prints what engine.attachments makes of them for each backend.

    set QT_QPA_PLATFORM=offscreen
    uv run python tools/attachments_test.py
"""

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from PySide6 import QtCore, QtGui, QtWidgets  # noqa: E402

from engine.attachments import blocks_for, kind_of  # noqa: E402

OUT = Path(__file__).resolve().parent.parent / "runtime" / "attach_samples"
OUT.mkdir(parents=True, exist_ok=True)
app = QtWidgets.QApplication(sys.argv)


def make_pdf(path: Path) -> None:
    w = QtGui.QPdfWriter(str(path))
    w.setPageSize(QtGui.QPageSize(QtGui.QPageSize.A4))
    p = QtGui.QPainter(w)
    p.setFont(QtGui.QFont("Arial", 11))
    y = 300
    for line in ["Bracket BR-200 specification", "Material: 6061-T6 aluminium",
                 "Base plate 120 x 60 x 8 mm, four M6 clearance holes on a 100 x 40 mm pattern.",
                 "Fillet all external edges R2. Anodise clear."] * 4:
        p.drawText(300, y, line)
        y += 400
    w.newPage()  # page 2: a drawing, no text
    p.setPen(QtGui.QPen(QtCore.Qt.black, 20))
    p.drawRect(1500, 2000, 6000, 3000)
    for x, yy in ((2300, 2700), (6700, 2700), (2300, 4300), (6700, 4300)):
        p.drawEllipse(QtCore.QPoint(x, yy), 250, 250)
    p.end()


def make_office() -> None:
    import docx
    import openpyxl
    import pptx
    d = docx.Document()
    d.add_heading("Design review notes", 1)
    d.add_paragraph("Increase wall thickness to 3 mm near the boss.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text, t.cell(1, 0).text, t.cell(1, 1).text = "Item", "Action", "Boss", "Add rib"
    d.save(OUT / "review.docx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "loads"
    for row in (("case", "Fx [N]", "Fy [N]"), ("LC1", 1200, 0), ("LC2", 0, 850)):
        ws.append(row)
    wb.save(OUT / "loads.xlsx")
    prs = pptx.Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Concept B"
    s.placeholders[1].text = "Lighter by 18 %, same stiffness"
    prs.save(OUT / "concepts.pptx")


make_pdf(OUT / "spec.pdf")
make_office()
(OUT / "points.csv").write_text("x,y,z\n0,0,0\n10,0,5\n", encoding="utf-8")
(OUT / "part.step").write_text("ISO-10303-21;\nHEADER;\nENDSEC;\n", encoding="utf-8")
img = QtGui.QImage(2400, 1600, QtGui.QImage.Format_RGB32)
img.fill(QtGui.QColor("#88a"))
img.save(str(OUT / "photo.png"))

files = [str(OUT / n) for n in ("spec.pdf", "review.docx", "loads.xlsx", "concepts.pptx",
                                 "points.csv", "part.step", "photo.png")]
print("kinds:", {Path(f).name: kind_of(Path(f)) for f in files})
for backend, budget in (("ollama", 60000), ("anthropic", 400000), ("ollama-tiny-budget", 200)):
    blocks = blocks_for(files, backend.split("-")[0], budget)
    print(f"\n== {backend} (budget {budget}) -> {len(blocks)} blocks")
    for b in blocks:
        if b["type"] == "text":
            print("  text :", b["text"][:150].replace("\n", " | "))
        elif b["type"] == "image":
            print(f"  image: {b['source']['media_type']} {len(b['source']['data']) * 3 // 4 // 1024} KB")
        else:
            print(f"  {b['type']}: {b.get('title')} {b['source']['media_type']}")
