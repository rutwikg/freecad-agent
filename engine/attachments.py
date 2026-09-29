"""Turn files the user drops into the chat into message content blocks.

Mirrors what the Claude desktop app does with attachments, per backend:
- Anthropic takes images and PDFs natively (a PDF keeps its visuals).
- A local model behind Ollama takes only text and images, so PDFs become page text
  plus rendered images of pages that are mostly drawing, and Office files become text.

Everything shares one character budget per message, because a local model's context
is small (64k tokens, summarised at ~32k). Over budget, content is truncated with a
note telling the agent where the full file is, so it can Read the rest on demand.
"""

from __future__ import annotations

import base64
import io
from pathlib import Path
from typing import Any

IMAGE_MEDIA = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
               ".webp": "image/webp", ".gif": "image/gif", ".bmp": "image/png"}
TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".rst", ".csv", ".tsv", ".json", ".jsonl", ".xml",
                 ".yaml", ".yml", ".toml", ".ini", ".cfg", ".conf", ".log", ".py", ".fcmacro",
                 ".js", ".ts", ".html", ".css", ".c", ".h", ".cpp", ".hpp", ".java", ".cs", ".m",
                 ".sql", ".sh", ".bat", ".ps1", ".tex", ".inp", ".dat", ".sta", ".msg", ".f", ".for"}
# Geometry files: not useful as text in a prompt; the agent opens them with CAD tools.
CAD_SUFFIXES = {".step", ".stp", ".iges", ".igs", ".stl", ".obj", ".fcstd", ".brep", ".brp",
                ".dxf", ".dwg", ".3mf", ".ply", ".off", ".catpart", ".catproduct", ".sldprt", ".x_t"}

MAX_IMAGE_EDGE = 1568  # Anthropic's recommended long edge; plenty for local vision models
MAX_PDF_IMAGES = 6
DRAWING_TEXT_CHARS = 300  # a page with less text than this is probably a drawing or a scan


def kind_of(path: Path) -> str:
    s = path.suffix.lower()
    if s in IMAGE_MEDIA:
        return "image"
    if s == ".pdf":
        return "pdf"
    if s in (".docx", ".xlsx", ".xlsm", ".pptx"):
        return "office"
    if s in CAD_SUFFIXES:
        return "cad"
    if s in TEXT_SUFFIXES or _looks_like_text(path):
        return "text"
    return "binary"


def _looks_like_text(path: Path) -> bool:
    try:
        chunk = path.read_bytes()[:4096]
    except OSError:
        return False
    return b"\0" not in chunk and _decode(chunk) is not None


def _decode(data: bytes) -> str | None:
    for enc in ("utf-8", "utf-16", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return None


def _image_block(img_bytes: bytes, media: str) -> dict[str, Any]:
    return {"type": "image", "source": {"type": "base64", "media_type": media,
                                        "data": base64.b64encode(img_bytes).decode("ascii")}}


def _fit_image(path_or_pil) -> tuple[bytes, str]:
    """Downscale to MAX_IMAGE_EDGE and re-encode; returns (bytes, media type)."""
    from PIL import Image
    img = path_or_pil if hasattr(path_or_pil, "size") else Image.open(path_or_pil)
    if max(img.size) > MAX_IMAGE_EDGE:
        img = img.copy()
        img.thumbnail((MAX_IMAGE_EDGE, MAX_IMAGE_EDGE))
    buf = io.BytesIO()
    if img.mode in ("RGBA", "LA", "P"):
        img.convert("RGBA").save(buf, "PNG")
        return buf.getvalue(), "image/png"
    img.convert("RGB").save(buf, "JPEG", quality=88)
    return buf.getvalue(), "image/jpeg"


class Budget:
    def __init__(self, chars: int):
        self.left = chars

    def take(self, text: str, path: Path) -> str:
        if len(text) <= self.left:
            self.left -= len(text)
            return text
        kept = text[: max(self.left, 0)]
        self.left = 0
        return (kept + f"\n\n[… truncated: {len(text) - len(kept):,} more characters. "
                f"The full file is at {path}; use the Read tool for the rest.]")


def _wrap(path: Path, body: str, note: str = "") -> dict[str, Any]:
    head = f'<attachment name="{path.name}" path="{path}"{" " + note if note else ""}>'
    return {"type": "text", "text": f"{head}\n{body}\n</attachment>"}


def _pdf_blocks(path: Path, backend: str, budget: Budget) -> list[dict[str, Any]]:
    if backend == "anthropic" and path.stat().st_size < 30 * 1024 * 1024:
        return [{"type": "document", "title": path.name, "source": {
            "type": "base64", "media_type": "application/pdf",
            "data": base64.b64encode(path.read_bytes()).decode("ascii")}}]
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    pages = [(p.extract_text() or "").strip() for p in reader.pages]
    text = "\n\n".join(f"--- page {i + 1} ---\n{t}" for i, t in enumerate(pages))
    blocks = [_wrap(path, budget.take(text, path), f'pages="{len(pages)}"')]
    # Drawings and scans carry little extractable text: send them as images instead.
    visual = [i for i, t in enumerate(pages) if len(t) < DRAWING_TEXT_CHARS][:MAX_PDF_IMAGES]
    if visual:
        import pypdfium2 as pdfium
        pdf = pdfium.PdfDocument(str(path))
        for i in visual:
            page = pdf[i]
            scale = MAX_IMAGE_EDGE / max(page.get_size())
            data, media = _fit_image(page.render(scale=scale).to_pil())
            blocks.append({"type": "text", "text": f"[{path.name}, page {i + 1} as an image:]"})
            blocks.append(_image_block(data, media))
        pdf.close()
    return blocks


def _office_text(path: Path) -> str:
    s = path.suffix.lower()
    if s == ".docx":
        import docx
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            parts.append("\n".join(" | ".join(c.text.strip() for c in row.cells) for row in t.rows))
        return "\n\n".join(parts)
    if s in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        out = []
        for ws in wb.worksheets:
            rows = [",".join("" if v is None else str(v) for v in row)
                    for row in ws.iter_rows(values_only=True)]
            out.append(f"--- sheet {ws.title} ---\n" + "\n".join(r for r in rows if r.strip(",")))
        return "\n\n".join(out)
    if s == ".pptx":
        import pptx
        prs = pptx.Presentation(str(path))
        out = []
        for i, slide in enumerate(prs.slides, 1):
            texts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
            out.append(f"--- slide {i} ---\n" + "\n".join(texts))
        return "\n\n".join(out)
    return ""


def blocks_for(paths: list[str], backend: str, char_budget: int) -> list[dict[str, Any]]:
    """Content blocks for all attachments of one message, within one shared budget."""
    budget = Budget(char_budget)
    blocks: list[dict[str, Any]] = []
    for p in paths:
        path = Path(p)
        try:
            kind = kind_of(path)
            if kind == "image":
                data, media = _fit_image(path)
                blocks.append({"type": "text", "text": f"[image: {path.name}]"})
                blocks.append(_image_block(data, media))
            elif kind == "pdf":
                blocks += _pdf_blocks(path, backend, budget)
            elif kind == "office":
                blocks.append(_wrap(path, budget.take(_office_text(path), path)))
            elif kind == "text":
                blocks.append(_wrap(path, budget.take(_decode(path.read_bytes()) or "", path)))
            else:  # CAD or unknown binary: point the agent at it rather than inlining bytes
                blocks.append({"type": "text", "text": f"[attached {kind} file: {path} — open or "
                                                       f"inspect it with the appropriate tools]"})
        except Exception as e:  # a bad file must not sink the whole message
            blocks.append({"type": "text", "text": f"[could not read attachment {path}: {e}]"})
    return blocks
